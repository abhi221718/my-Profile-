import json
import os
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pymongo import MongoClient
from pymongo.errors import PyMongoError
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_DOCX_ARCHIVE_ENTRIES = 2000
MAX_JOB_DESCRIPTION_SIZE = 12_000
MAX_JOBS = 10
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

load_dotenv(Path(__file__).resolve().parents[1] / ".env")

MONGODB_URI = os.getenv("MONGODB_URI", "").strip()
FRONTEND_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        "FRONTEND_ORIGINS",
        "http://localhost:5173,http://127.0.0.1:5173,http://localhost:4173,http://127.0.0.1:4173",
    ).split(",")
    if origin.strip()
]

SKILL_ALIASES = {
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure", "microsoft azure"],
    "C": ["c"],
    "C#": ["c#"],
    "C++": ["c++"],
    "CSS": ["css", "css3"],
    "Django": ["django"],
    "Docker": ["docker", "containerization"],
    "Figma": ["figma"],
    "Flask": ["flask"],
    "Git": ["git", "github", "gitlab", "version control"],
    "HTML": ["html", "html5"],
    "Java": ["java"],
    "JavaScript": ["javascript", "ecmascript"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Machine Learning": ["machine learning", "ml"],
    "MongoDB": ["mongodb", "mongo db"],
    "Natural Language Processing": ["natural language processing", "nlp"],
    "Next.js": ["next.js", "nextjs"],
    "Node.js": ["node.js", "nodejs"],
    "NumPy": ["numpy"],
    "Pandas": ["pandas"],
    "PHP": ["php"],
    "PostgreSQL": ["postgresql", "postgres"],
    "Power BI": ["power bi", "powerbi"],
    "Python": ["python"],
    "PyTorch": ["pytorch"],
    "React": ["react", "react.js", "reactjs"],
    "REST API": ["rest api", "restful api", "restful services"],
    "SQL": ["sql", "mysql", "sqlite", "sql server"],
    "Scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "Spark": ["apache spark", "pyspark"],
    "Tableau": ["tableau"],
    "TensorFlow": ["tensorflow", "keras"],
    "TypeScript": ["typescript"],
    "UI/UX Design": ["ui/ux", "ui design", "ux design", "user experience design"],
    "Vue.js": ["vue.js", "vuejs"],
    "WordPress": ["wordpress"],
    "Communication": ["communication", "interpersonal skills"],
    "Data Analysis": ["data analysis", "data analytics", "data analyst"],
    "Data Structures": ["data structures", "dsa", "algorithms"],
    "Excel": ["excel", "microsoft excel"],
    "Express.js": ["express.js", "expressjs"],
    "FastAPI": ["fastapi", "fast api"],
    "Jupyter": ["jupyter", "jupyter notebook"],
    "Linux": ["linux", "unix"],
    "Matplotlib": ["matplotlib"],
    "OOP": ["object-oriented programming", "object oriented programming", "oop"],
    "Problem Solving": ["problem solving", "problem-solving"],
    "Responsive Design": ["responsive design", "mobile-first design"],
    "Sass": ["sass", "scss"],
    "Seaborn": ["seaborn"],
    "Software Testing": ["software testing", "unit testing", "pytest", "jest"],
    "Statistics": ["statistics", "statistical analysis"],
    "Teamwork": ["teamwork", "collaboration"],
    "Vite": ["vite"],
}

EDUCATION_TERMS = (
    "bachelor",
    "master",
    "b.tech",
    "btech",
    "m.tech",
    "mtech",
    "b.e.",
    "b.sc",
    "m.sc",
    "ph.d",
    "diploma",
    "university",
    "college",
    "education",
    "degree",
    "bca",
    "mca",
)
EXPERIENCE_TERMS = (
    "experience",
    "employment",
    "work history",
    "internship",
    "intern",
    "years",
)

app = FastAPI(
    title="AI Resume Analyzer & Job Matcher",
    description="Extract resume details, identify skills, and rank job matches using TF-IDF and cosine similarity.",
    version="1.0.0",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)


def extract_pdf(data: bytes) -> str:
    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.is_encrypted:
                raise HTTPException(
                    status_code=422,
                    detail="This PDF is password-protected. Upload an unlocked copy.",
                )
            if len(document) > 100:
                raise HTTPException(
                    status_code=422,
                    detail="This PDF has more than 100 pages. Upload a shorter resume.",
                )
            return "\n".join(page.get_text() for page in document)
    except HTTPException:
        raise
    except (fitz.FileDataError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not read this PDF. Check that the file is valid.",
        ) from exc


def extract_docx(data: bytes) -> str:
    try:
        with ZipFile(BytesIO(data)) as package:
            entries = package.infolist()
            if len(entries) > MAX_DOCX_ARCHIVE_ENTRIES or sum(
                entry.file_size for entry in entries
            ) > MAX_DOCX_UNCOMPRESSED_SIZE:
                raise HTTPException(
                    status_code=413,
                    detail="This DOCX expands beyond the supported size. Upload a smaller resume.",
                )
            if "word/document.xml" not in package.namelist():
                raise HTTPException(
                    status_code=422,
                    detail="This file is not a valid DOCX resume.",
                )
        document = Document(BytesIO(data))
    except HTTPException:
        raise
    except (BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not read this DOCX. Check that the file is valid.",
        ) from exc

    sections = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                sections.append(" | ".join(cells))
    return "\n".join(sections)


def clean_text(text: str) -> str:
    cleaned_lines = []
    for line in text.replace("\x00", " ").splitlines():
        normalized = re.sub(r"\s+", " ", line).strip()
        if normalized:
            cleaned_lines.append(normalized)
    return "\n".join(cleaned_lines)


def contains_term(text: str, term: str) -> bool:
    pattern = rf"(?<![a-z0-9]){re.escape(term.lower())}(?![a-z0-9])"
    return re.search(pattern, text.lower()) is not None


def extract_skills(text: str) -> list[str]:
    normalized = text.lower()
    return [
        skill
        for skill, aliases in SKILL_ALIASES.items()
        if any(
            re.search(r"(?<![a-z0-9+#])c(?![a-z0-9+#])", normalized)
            if skill == "C" and alias == "c"
            else contains_term(normalized, alias)
            for alias in aliases
        )
    ]


def extract_candidate(text: str) -> dict[str, str]:
    email_match = re.search(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.IGNORECASE
    )
    phone_match = re.search(
        r"(?<!\w)(?:\+?\d{1,3}[\s().-]?)?(?:\d[\s().-]?){8,13}\d(?!\w)",
        text,
    )
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    name = ""
    non_name_terms = (
        "resume",
        "curriculum vitae",
        "engineer",
        "developer",
        "analyst",
        "manager",
        "profile",
        "contact",
        "skills",
        "experience",
        "education",
        "portfolio",
        "summary",
        "objective",
        "student",
    )
    for line in lines[:8]:
        if (
            "@" in line
            or re.search(r"\d", line)
            or len(line) > 70
            or any(contains_term(line, term) for term in non_name_terms)
        ):
            continue
        words = re.findall(r"[A-Za-z][A-Za-z'’-]*", line)
        if 2 <= len(words) <= 4 and len(line.split()) <= 5:
            name = line
            break
    return {
        "name": name,
        "email": email_match.group(0) if email_match else "",
        "phone": re.sub(r"\s+", " ", phone_match.group(0)).strip() if phone_match else "",
    }


def extract_section_lines(text: str, terms: tuple[str, ...]) -> list[str]:
    matches = []
    for line in text.splitlines():
        normalized = line.lower()
        if any(contains_term(normalized, term) for term in terms):
            matches.append(line[:180])
        if len(matches) == 5:
            break
    return matches


def has_quantified_achievement(text: str) -> bool:
    metric_pattern = (
        r"\b\d+(?:\.\d+)?\s*(?:%|\b(?:users?|customers?|clients?|orders?|requests?|"
        r"tickets?|hours?|minutes?|seconds?|workflows?|transactions?)\b)|"
        r"[$€£]\s?\d[\d,.]*"
    )
    impact_pattern = (
        r"\b(?:increased|improved|reduced|saved|generated|delivered|grew|"
        r"cut|boosted|achieved)\b[^.\n]{0,60}\b\d+(?:\.\d+)?\b"
    )
    return re.search(metric_pattern, text, re.IGNORECASE) is not None or re.search(
        impact_pattern, text, re.IGNORECASE
    ) is not None


def calculate_resume_score(text: str, candidate: dict[str, str], skills: list[str]) -> tuple[int, dict[str, int]]:
    normalized = text.lower()
    breakdown = {
        "contact_details": (5 if candidate["name"] else 0)
        + (7 if candidate["email"] else 0)
        + (3 if candidate["phone"] else 0),
        "skills_section": 20 if contains_term(normalized, "skills") else 0,
        "education": 15
        if any(contains_term(normalized, term) for term in EDUCATION_TERMS)
        else 0,
        "experience": 15
        if any(contains_term(normalized, term) for term in EXPERIENCE_TERMS)
        else 0,
        "projects": 10 if contains_term(normalized, "project") else 0,
        "relevant_skills": min(len(skills) * 3, 15),
        "measurable_achievements": 10 if has_quantified_achievement(text) else 0,
    }
    return sum(breakdown.values()), breakdown


def compare_job(resume_text: str, resume_skills: list[str], job: dict) -> dict:
    title = job.get("title", "").strip()[:100] or "Untitled job"
    description = job["description"].strip()
    job_skills = extract_skills(description)
    matched_skills = [skill for skill in job_skills if skill in resume_skills]
    missing_skills = [skill for skill in job_skills if skill not in resume_skills]

    try:
        vectors = TfidfVectorizer(stop_words="english", ngram_range=(1, 2)).fit_transform(
            [resume_text, description]
        )
        text_similarity = float(cosine_similarity(vectors[0:1], vectors[1:2])[0][0])
    except ValueError:
        text_similarity = 0.0

    if job_skills:
        skill_coverage = len(matched_skills) / len(job_skills)
        match_score = (skill_coverage * 0.65) + (text_similarity * 0.35)
    else:
        skill_coverage = 0.0
        match_score = text_similarity

    return {
        "title": title,
        "match_percentage": round(match_score * 100),
        "text_similarity": round(text_similarity * 100),
        "skill_coverage": round(skill_coverage * 100),
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "required_skills": job_skills,
    }


def make_suggestions(text: str, candidate: dict[str, str], skills: list[str], matches: list[dict]) -> list[str]:
    normalized = text.lower()
    suggestions = []
    if not candidate["email"] or not candidate["phone"]:
        suggestions.append("Add a clearly labeled email address and phone number to your contact section.")
    if not any(contains_term(normalized, term) for term in EXPERIENCE_TERMS):
        suggestions.append("Add a work experience or internship section with role titles, organizations, and dates.")
    if not any(contains_term(normalized, term) for term in EDUCATION_TERMS):
        suggestions.append("Add your education, qualification, institution, and graduation year.")
    if not contains_term(normalized, "project"):
        suggestions.append("Include relevant projects and describe what you built and which tools you used.")
    if not has_quantified_achievement(text):
        suggestions.append("Where accurate, quantify results with metrics such as time saved, users served, or performance improved.")
    if len(skills) < 5:
        suggestions.append("Add a concise, clearly labeled skills section using the tools and technologies you genuinely know.")

    prioritized = sorted(matches, key=lambda match: match["match_percentage"], reverse=True)
    for match in prioritized:
        for skill in match["missing_skills"][:2]:
            suggestions.append(
                f"Build hands-on experience with {skill}, then add it to your resume when you can support it with an example."
            )
        if len(suggestions) >= 8:
            break
    return suggestions[:8]


def save_analysis(record: dict) -> tuple[str, str]:
    if not MONGODB_URI:
        return "disabled", ""
    try:
        with MongoClient(MONGODB_URI, serverSelectionTimeoutMS=2500) as client:
            database_name = os.getenv("MONGODB_DATABASE", "ai_resume_analyzer")
            collection = client[database_name]["analyses"]
            collection.insert_one(record)
        return "saved", ""
    except PyMongoError:
        return (
            "error",
            "MongoDB connection or write failed. Check the API host logs and connection settings.",
        )


@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "storage": "configured" if MONGODB_URI else "disabled",
    }


@app.post("/api/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    jobs: str = Form(...),
):
    extension = Path(file.filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Only PDF and DOCX resumes are supported.")

    contents = await file.read(MAX_FILE_SIZE + 1)
    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded resume is empty.")
    if len(contents) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Resume file is larger than 10 MB.")

    try:
        job_items = json.loads(jobs)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Job descriptions must be valid JSON.") from exc
    if not isinstance(job_items, list) or not job_items:
        raise HTTPException(status_code=400, detail="Add at least one job description.")
    if len(job_items) > MAX_JOBS:
        raise HTTPException(status_code=400, detail=f"You can compare up to {MAX_JOBS} jobs at a time.")

    normalized_jobs = []
    for item in job_items:
        if not isinstance(item, dict) or not isinstance(item.get("description"), str):
            raise HTTPException(status_code=400, detail="Each job needs a description.")
        description = item["description"].strip()
        if not description:
            continue
        if len(description) > MAX_JOB_DESCRIPTION_SIZE:
            raise HTTPException(
                status_code=413,
                detail=f"Each job description must be {MAX_JOB_DESCRIPTION_SIZE} characters or fewer.",
            )
        normalized_jobs.append(
            {"title": str(item.get("title", "")), "description": description}
        )
    if not normalized_jobs:
        raise HTTPException(status_code=400, detail="Add at least one non-empty job description.")

    extracted = extract_pdf(contents) if extension == ".pdf" else extract_docx(contents)
    text = clean_text(extracted)
    if not text:
        raise HTTPException(
            status_code=422,
            detail="No selectable text was found. Scanned/image-only PDFs need OCR before analysis.",
        )

    candidate = extract_candidate(text)
    skills = extract_skills(text)
    education = extract_section_lines(text, EDUCATION_TERMS)
    experience = extract_section_lines(text, EXPERIENCE_TERMS)
    resume_score, score_breakdown = calculate_resume_score(text, candidate, skills)
    job_matches = [
        compare_job(text, skills, job)
        for job in normalized_jobs
    ]
    job_matches.sort(key=lambda match: match["match_percentage"], reverse=True)
    suggestions = make_suggestions(text, candidate, skills, job_matches)

    response = {
        "filename": Path(file.filename or "resume").name,
        "candidate": candidate,
        "skills": skills,
        "education": education,
        "experience": experience,
        "resume_score": resume_score,
        "score_breakdown": score_breakdown,
        "job_matches": job_matches,
        "suggestions": suggestions,
    }
    record = {
        "created_at": datetime.now(timezone.utc),
        "filename": response["filename"],
        "candidate_name": candidate["name"],
        "resume_score": resume_score,
        "skills": skills,
        "job_matches": [
            {
                "title": match["title"],
                "match_percentage": match["match_percentage"],
            }
            for match in job_matches
        ],
    }
    storage_status, storage_message = save_analysis(record)
    response["storage_status"] = storage_status
    if storage_message:
        response["storage_message"] = storage_message
    return response


@app.get("/api/history")
def analysis_history():
    if not MONGODB_URI:
        raise HTTPException(
            status_code=503,
            detail="Analysis history is disabled. Configure MONGODB_URI to enable MongoDB storage.",
        )
    try:
        with MongoClient(MONGODB_URI, serverSelectionTimeoutMS=2500) as client:
            database_name = os.getenv("MONGODB_DATABASE", "ai_resume_analyzer")
            records = list(
                client[database_name]["analyses"]
                .find({}, {"_id": 0})
                .sort("created_at", -1)
                .limit(20)
            )
    except PyMongoError as exc:
        raise HTTPException(
            status_code=503,
            detail="Could not retrieve MongoDB analysis history. Check the API host logs and connection settings.",
        ) from exc

    for record in records:
        created_at = record.get("created_at")
        if created_at:
            record["created_at"] = created_at.isoformat()
    return {"items": records}
